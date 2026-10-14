import io, os, json, time, uuid, re, urllib.request, secrets, hmac, copy
from html import escape
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request
from fastapi.responses import StreamingResponse, HTMLResponse, FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from dotenv import load_dotenv
import pdfplumber, docx
from anthropic import Anthropic
from pptx import Presentation
from pptx.util import Inches, Pt

APP_DIR = os.path.dirname(os.path.abspath(__file__))


def load_environment():
    env_path = os.path.join(APP_DIR, ".env")
    if os.path.exists(env_path):
        load_dotenv(dotenv_path=env_path, override=False)
    return os.environ


load_environment()

app = FastAPI(title="Resume Presenter")
API_KEY = (os.getenv("ANTHROPIC_API_KEY") or "").strip()
WORKSPACE_ID = (os.getenv("ANTHROPIC_WORKSPACE_ID") or "").strip()
MODEL = os.getenv("MODEL", "claude-sonnet-5")
PUBLIC_BASE_URL = (os.getenv("PUBLIC_BASE_URL") or "https://sassbyte.site").strip().rstrip("/")
MAX_BYTES, MAX_CHARS, TTL = 5 * 1024 * 1024, 20000, 24 * 3600
PERSIST_BASE = os.path.expanduser("~/persist/storyboard")
os.makedirs(PERSIST_BASE, exist_ok=True)
DECKS: dict = {}  # id -> (expires_at, deck). Resume files are never written to disk.


def valid_deck_id(did: str) -> bool:
    return bool(re.fullmatch(r"[A-Za-z0-9_-]{1,100}", did or ""))


def persist_edit_token(did: str, token: str):
    if not valid_deck_id(did):
        raise ValueError("Invalid storyboard id")
    deck_dir = os.path.join(PERSIST_BASE, did)
    os.makedirs(deck_dir, mode=0o700, exist_ok=True)
    token_path = os.path.join(deck_dir, "owner.token")
    fd = os.open(token_path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, "w", encoding="utf-8") as token_file:
        token_file.write(token)
    os.chmod(token_path, 0o600)


def can_edit_deck(did: str, token: str) -> bool:
    if not token or not valid_deck_id(did):
        return False
    token_path = os.path.join(PERSIST_BASE, did, "owner.token")
    try:
        with open(token_path, "r", encoding="utf-8") as token_file:
            saved_token = token_file.read()
    except FileNotFoundError:
        return False
    return hmac.compare_digest(saved_token, token)


def render_share_page(did: str, deck: dict) -> str:
    title = escape((deck.get("name") or "Resume Storyboard").strip() or "Resume Storyboard", quote=True)
    description = "Career story and presentation deck created from a resume."
    share_url = f"{PUBLIC_BASE_URL}/share/{did}"
    share_image = f"{PUBLIC_BASE_URL}/share-preview.svg"
    css = '''
    body { font-family: Inter, system-ui, sans-serif; margin: 0; background: #f5f7ff; color: #171a2d; }
    .wrap { max-width: 980px; margin: 0 auto; padding: 40px 18px 60px; }
    .panel { background: #fff; border: 1px solid #dfe6ff; border-radius: 22px; box-shadow: 0 18px 40px rgba(35,42,87,.12); padding: 28px; }
    h1 { margin: 0 0 10px; font-size: clamp(2rem,4vw,3rem); letter-spacing: -.05em; }
    .sub { color: #5f6888; line-height: 1.7; }
    .slide { margin-top: 18px; border: 1px solid #dde4ff; border-radius: 18px; padding: 22px; background: #ffffff; }
    .slide h2 { margin: 0 0 18px; font-size: clamp(1.3rem,3vw,2rem); color: #2f35dd; }
    .item { margin: 12px 0; padding-left: 16px; border-left: 3px solid #ecf0ff; }
    .item b { display: block; margin-bottom: 4px; }
    .bar { background: #eef0ff; border-radius: 999px; height: 12px; overflow: hidden; margin-top: 6px; }
    .bar i { display: block; height: 100%; width: 0; background: linear-gradient(90deg,#4549f5,#7b8bff); border-radius: 999px; }
    .stats { display: grid; grid-template-columns: repeat(auto-fit,minmax(130px,1fr)); gap: 12px; }
    .stat { background: #eef0ff; border-radius: 16px; padding: 16px; text-align: center; border: 1px solid #dde4ff; }
    .stat b { display: block; color: #2f35dd; font-size: 2rem; }
    .nav { display: flex; gap: 10px; justify-content: center; flex-wrap: wrap; margin-top: 20px; }
    button { border: 0; border-radius: 12px; padding: 10px 18px; background: #4549f5; color: white; cursor: pointer; }
    button:disabled { opacity: .55; cursor: not-allowed; }
    .muted { color: #5f6888; font-size: 14px; }
    @media (max-width: 640px) { .slide { padding: 18px; } }
    '''
    js = f'''
    const deckId = "{did}";
    let deck = null;
    let index = 0;
    let autoplayTimer = null;
    let narrationTimer = null;
    let narrationActive = false;
    const deckNode = document.getElementById('deck');
    const playButton = document.getElementById('play');
    const stopButton = document.getElementById('stopAudio');
    const narrationStatus = document.getElementById('narrationStatus');
    const params = new URLSearchParams(window.location.search);
    const autoplay = params.get('autoplay') === '1';
    const userFactIds = new Set({json.dumps([
        fact.get("id") for fact in (deck.get("profile", {}).get("facts", []) if isinstance(deck.get("profile"), dict) else [])
        if isinstance(fact, dict) and fact.get("source") == "user_chat"
        and isinstance(fact.get("id"), str) and len(fact["id"]) == 32
        and all(char in "abcdef0123456789" for char in fact["id"])
    ])});

    function escapeHtml(value) {{
      return String(value ?? '').replace(/[&<>"']/g, (char) => ({{
        '&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;'
      }})[char]);
    }}

    function startAutoplay() {{
      if (!deck || deck.slides.length < 2) return;
      stopAutoplay();
      autoplayTimer = setInterval(() => {{
        if (index >= deck.slides.length - 1) {{
          index = 0;
        }} else {{
          index += 1;
        }}
        render();
      }}, 7000);
    }}

    function stopAutoplay() {{
      if (autoplayTimer) clearInterval(autoplayTimer);
      autoplayTimer = null;
    }}

        function stopNarration() {{
            narrationActive = false;
            if (narrationTimer) clearTimeout(narrationTimer);
            narrationTimer = null;
            if ('speechSynthesis' in window) window.speechSynthesis.cancel();
            playButton.textContent = 'Play narration';
            stopButton.disabled = true;
        }}

        function narrationText(slide) {{
            const pronouns = {{ F: ['she', 'her', 'She'], M: ['he', 'his', 'He'], N: ['they', 'their', 'They'] }};
            const [she, her, Her] = pronouns[deck.gender] || pronouns.N;
            const fallback = [slide.h, slide.s, (slide.items || []).flat().join('. ')].filter(Boolean).join('. ');
            return (slide.n || fallback)
                .replace(/\\{{name\\}}/g, deck.name || '')
                .replace(/\\{{she\\}}/g, she)
                .replace(/\\{{her\\}}/g, her)
                .replace(/\\{{Her\\}}/g, Her);
        }}

        function speakSlide() {{
            if (!narrationActive || !deck) return;
            const utterance = new SpeechSynthesisUtterance(narrationText(deck.slides[index]));
            const voice = window.speechSynthesis.getVoices().find((item) => /^en/i.test(item.lang));
            if (voice) utterance.voice = voice;
            utterance.onend = () => {{
                if (!narrationActive) return;
                if (index < deck.slides.length - 1) {{
                    index += 1;
                    render();
                    narrationTimer = setTimeout(speakSlide, 350);
                }} else {{
                    stopNarration();
                    narrationStatus.textContent = 'Narration complete';
                }}
            }};
            utterance.onerror = () => stopNarration();
            window.speechSynthesis.speak(utterance);
        }}

        function toggleNarration() {{
            if (!('speechSynthesis' in window) || typeof SpeechSynthesisUtterance === 'undefined') {{
                narrationStatus.textContent = 'Speech playback is not supported in this browser.';
                return;
            }}
            if (narrationActive && window.speechSynthesis.paused) {{
                window.speechSynthesis.resume();
                playButton.textContent = 'Pause narration';
                return;
            }}
            if (narrationActive && window.speechSynthesis.speaking) {{
                window.speechSynthesis.pause();
                playButton.textContent = 'Resume narration';
                return;
            }}
            stopAutoplay();
            narrationActive = true;
            narrationStatus.textContent = '';
            playButton.textContent = 'Pause narration';
            stopButton.disabled = false;
            window.speechSynthesis.cancel();
            speakSlide();
        }}

    function render() {{
      const slide = deck.slides[index];
      if (!slide) return;
      const addedBadge = Array.isArray(slide.fact_ids) && slide.fact_ids.some((id) => userFactIds.has(id))
        ? '<span class="muted"> · Added by candidate</span>' : '';
      let html = '';
      if (slide.t === 'title') {{
        html = '<div class="slide" style="padding:52px 26px;text-align:center"><h2 style="margin-bottom:10px">' + escapeHtml(slide.h) + '</h2><div class="sub">' + escapeHtml(slide.s || '') + '</div>' + addedBadge + '</div>';
      }} else {{
        html = '<div class="slide"><h2>' + escapeHtml(slide.h) + '</h2>' + addedBadge;
        const items = slide.items || [];
        if (slide.t === 'timeline') {{
          html += items.map((x) => '<div class="item"><b>' + escapeHtml(x[0]) + '</b><span>' + escapeHtml(x[1]) + '</span></div>').join('');
        }}
        if (slide.t === 'list') {{
          html += items.map((x) => '<div class="item"><b>' + escapeHtml(x[0]) + '</b><span>' + escapeHtml(x[1]) + '</span></div>').join('');
        }}
        if (slide.t === 'skills') {{
          html += items.map((x) => '<div class="item"><b>' + escapeHtml(x[0]) + '</b><div class="bar"><i data-w="' + escapeHtml(x[1]) + '\"></i></div></div>').join('');
        }}
        if (slide.t === 'stats') {{
          html += '<div class="stats">' + items.map((x) => '<div class="stat"><b>' + escapeHtml(x[0]) + '</b><div>' + escapeHtml(x[1]) + '</div></div>').join('') + '</div>';
        }}
        html += '</div>';
      }}
      deckNode.innerHTML = html;
      document.querySelectorAll('[data-w]').forEach((bar) => {{
        bar.style.width = bar.dataset.w + '%';
      }});
    }}

    document.getElementById('next').onclick = () => {{
      if (!deck) return;
      stopAutoplay();
            stopNarration();
      index = Math.min(index + 1, deck.slides.length - 1);
      render();
      if (autoplay) startAutoplay();
    }};

    document.getElementById('prev').onclick = () => {{
      if (!deck) return;
      stopAutoplay();
            stopNarration();
      index = Math.max(index - 1, 0);
      render();
      if (autoplay) startAutoplay();
    }};

    fetch('/api/share/' + deckId).then((r) => r.json()).then((data) => {{
      deck = data;
      render();
    document.getElementById('play').onclick = toggleNarration;
    document.getElementById('stopAudio').onclick = stopNarration;
      if (autoplay) startAutoplay();
    }}).catch(() => {{
      deckNode.innerHTML = '<div class="slide"><h2>Storyboard unavailable</h2><div class="sub">This shared deck could not be loaded.</div></div>';
    }});
    '''
    return (
        '<!DOCTYPE html><html lang="en"><head>'
        '<meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        f'<title>{title} | Resume Storyboard</title>'
        f'<meta property="og:title" content="{title} | Resume Storyboard">'
        f'<meta property="og:description" content="{description}">'
        '<meta property="og:type" content="website">'
        f'<meta property="og:url" content="{share_url}">'
        '<meta property="og:site_name" content="Sassbyte">'
        f'<meta property="og:image" content="{share_image}">'
        f'<meta property="og:image:secure_url" content="{share_image}">'
        '<meta property="og:image:type" content="image/svg+xml">'
        '<meta property="og:image:width" content="1200">'
        '<meta property="og:image:height" content="630">'
        f'<meta property="og:image:alt" content="{title} Resume Storyboard preview">'
        '<meta name="twitter:card" content="summary_large_image">'
        f'<meta name="twitter:title" content="{title} | Resume Storyboard">'
        f'<meta name="twitter:description" content="{description}">'
        f'<meta name="twitter:image" content="{share_image}">'
        f'<link rel="canonical" href="{share_url}">'
        f'<style>{css}</style>'
        '</head><body>'
        '<div class="wrap"><div class="panel"><h1>' + title + '</h1><div class="sub">Resume storyboard shared for playback and presentation.</div><div id="deck"></div><div class="nav"><button id="prev">Prev</button><button id="play">Play narration</button><button id="stopAudio" disabled>Stop</button><button id="next">Next</button></div><div id="narrationStatus" class="muted" aria-live="polite" style="text-align:center;margin-top:10px"></div></div></div>'
        '<script>' + js + '</script>'
        '</body></html>'
    )
def persist_storyboard(did: str, deck: dict) -> str:
    deck_dir = os.path.join(PERSIST_BASE, str(did))
    os.makedirs(deck_dir, exist_ok=True)
    with open(os.path.join(deck_dir, "deck.json"), "w", encoding="utf-8") as f:
        json.dump(deck, f, ensure_ascii=False)
    with open(os.path.join(deck_dir, "index.html"), "w", encoding="utf-8") as f:
        f.write(render_share_page(did, deck))
    return deck_dir


def load_persisted_storyboard(did: str):
    if not valid_deck_id(did):
        return None
    deck_path = os.path.join(PERSIST_BASE, str(did), "deck.json")
    if not os.path.exists(deck_path):
        return None
    with open(deck_path, "r", encoding="utf-8") as f:
        return json.load(f)


def get_live_deck(did: str):
    if not valid_deck_id(did):
        return None
    now = time.time()
    if did in DECKS:
        expires_at, deck = DECKS[did]
        if expires_at > now:
            return deck
        DECKS.pop(did, None)

    deck = load_persisted_storyboard(did)
    if deck is None:
        return None

    DECKS[did] = (now + TTL, deck)
    return deck


def sanitize_error_message(message: str) -> str:
    text = (message or "").strip()
    lower = text.lower()

    if any(token in lower for token in [
        "invalid x-api-key",
        "x-api-key",
        "authentication_error",
        "invalid api key",
        "api key is invalid",
        "authentication",
    ]):
        return "Your Anthropic API key is invalid or expired. Please update it and try again."

    if any(token in lower for token in [
        "workspace",
        "billing",
        "credit",
        "quota",
        "rate limit",
        "usage",
        "not allowed",
    ]):
        return "Anthropic rejected this request. Please check your API key, workspace ID, and billing status."

    if "error code:" in lower or "traceback" in lower or "exception" in lower:
        return "Storyboard generation failed. Please check your Anthropic credentials and try again."

    return text or "Storyboard generation failed. Please check your Anthropic credentials and try again."


def sanitize_error_response(message: str):
    text = (message or "").strip()
    lower = text.lower()

    if any(token in lower for token in [
        "invalid x-api-key",
        "x-api-key",
        "authentication_error",
        "invalid api key",
        "api key is invalid",
        "authentication",
    ]):
        return 401, "Your Anthropic API key is invalid or expired. Please update it and try again."

    if any(token in lower for token in [
        "workspace",
        "billing",
        "credit",
        "quota",
        "rate limit",
        "usage",
        "not allowed",
    ]):
        return 400, "Anthropic rejected this request. Please check your API key, workspace ID, and billing status."

    return 502, "Storyboard generation failed. Please check your Anthropic credentials and try again."


def resolve_credentials(api_key: str = "", workspace_id: str = "", use_env: bool = False):
    if use_env:
        selected_key = (API_KEY or "").strip()
        selected_workspace = (WORKSPACE_ID or "").strip()
    else:
        selected_key = (api_key or "").strip()
        selected_workspace = (workspace_id or "").strip()
        if not selected_key:
            selected_key = (API_KEY or "").strip()
        if not selected_workspace:
            selected_workspace = (WORKSPACE_ID or "").strip()

    if selected_workspace and selected_workspace.lower().startswith("your_"):
        selected_workspace = ""
    return selected_key, selected_workspace


def get_client(api_key: str = "", workspace_id: str = "", use_env: bool = False):
    selected_key, selected_workspace = resolve_credentials(api_key, workspace_id, use_env)
    if not selected_key:
        return None, "Anthropic API key is missing. Enter your key in the form, tick 'Use rp/.env credentials', or add ANTHROPIC_API_KEY to rp/.env."
    headers = {"anthropic-version": "2023-06-01"}
    if selected_workspace:
        headers["anthropic-workspace-id"] = selected_workspace
    return Anthropic(api_key=selected_key, default_headers=headers), None

@app.get("/api/health")
def health():
    return {"status": "ok", "service": "resume-presenter"}


@app.post("/api/validate-key")
async def validate_key(request: Request):
    try:
        payload = await request.json()
    except Exception:
        payload = {}

    api_key = (payload.get("api_key") or "").strip()
    if not api_key:
        return JSONResponse(status_code=400, content={"valid": False, "detail": "Anthropic API key is required."})

    try:
        client = Anthropic(api_key=api_key)
        client.models.list()
        return {"valid": True, "detail": "Anthropic key validated successfully."}
    except Exception as exc:
        return JSONResponse(status_code=400, content={"valid": False, "detail": sanitize_error_message(str(exc))})


def get_cashfree_config():
    app_id = (os.getenv("CASHFREE_APP_ID") or "").strip()
    secret_key = (os.getenv("CASHFREE_SECRET_KEY") or "").strip()
    mode = (os.getenv("CASHFREE_MODE") or "sandbox").strip().lower()
    if mode not in {"sandbox", "production"}:
        mode = "sandbox"
    base_url = "https://sandbox.cashfree.com/pg" if mode == "sandbox" else "https://api.cashfree.com/pg"
    return {"app_id": app_id, "secret_key": secret_key, "mode": mode, "base_url": base_url}


@app.post("/api/payment/cashfree/create-session")
async def create_cashfree_session(request: Request):
    cfg = get_cashfree_config()
    if not cfg["app_id"] or not cfg["secret_key"]:
        raise HTTPException(400, "Cashfree is not configured. Set CASHFREE_APP_ID and CASHFREE_SECRET_KEY in the backend .env file.")

    try:
        payload = await request.json()
    except Exception:
        payload = {}

    amount = 50.0
    customer_name = (payload.get("customer_name") or "Sassbyte User").strip() or "Sassbyte User"
    customer_email = (payload.get("customer_email") or "user@sassbyte.site").strip() or "user@sassbyte.site"
    customer_phone = (payload.get("customer_phone") or "9999999999").strip() or "9999999999"
    order_id = f"sassbyte-{uuid.uuid4().hex[:12]}"
    host = request.headers.get("x-forwarded-host") or request.headers.get("host") or "localhost:8000"
    proto = request.headers.get("x-forwarded-proto") or request.url.scheme
    return_url = (payload.get("return_url") or f"{proto}://{host}/?payment=success").strip()
    notify_url = (payload.get("notify_url") or f"{proto}://{host}/api/payment/cashfree/notify").strip()

    order_payload = {
        "order_id": order_id,
        "order_amount": round(amount, 2),
        "order_currency": "INR",
        "customer_details": {
            "customer_id": f"customer-{uuid.uuid4().hex[:10]}",
            "customer_name": customer_name,
            "customer_email": customer_email,
            "customer_phone": customer_phone,
        },
        "order_meta": {
            "return_url": return_url,
            "notify_url": notify_url,
        },
    }

    try:
        req = urllib.request.Request(
            f"{cfg['base_url']}/orders",
            data=json.dumps(order_payload).encode("utf-8"),
            headers={
                "accept": "application/json",
                "content-type": "application/json",
                "x-client-id": cfg["app_id"],
                "x-client-secret": cfg["secret_key"],
            },
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=25) as resp:
            body = resp.read().decode("utf-8")
            data = json.loads(body or "{}")
    except Exception as exc:
        raise HTTPException(502, f"Cashfree payment initialization failed: {exc}") from exc

    payment_session_id = data.get("payment_session_id") or data.get("payment_session_id")
    if not payment_session_id:
        raise HTTPException(502, "Cashfree did not return a payment session id.")

    return {"success": True, "order_id": order_id, "payment_session_id": payment_session_id, "amount": round(amount, 2), "currency": "INR"}


@app.post("/api/payment/cashfree/notify")
async def cashfree_notify(request: Request):
    data = await request.json()
    return {"status": "received", "payload": data}

SYSTEM = """You turn a resume into a presentation storyboard. Return ONLY JSON, no markdown.
Rules: use ONLY facts in the resume. Never invent employers, dates, numbers or skills.
If a number is not in the resume, do not create one. Skill levels (0-100) are rough estimates from evidence.
Critical experience requirement: under the career journey / timeline slide, each experience must include the work performed, responsibilities, scope, tools, and relevant impact from the resume, not just the role and company name.
For each timeline item, format it as [role, "company · dates — 1-2 short but concrete work-performed details from the resume"].
Examples: ["Senior Product Manager", "Acme Inc · 2021-2024 — Led roadmap planning, partnered with engineering, and launched AI workflows for customer onboarding."]
Schema:
{"name":str,"gender":"F"|"M"|"N" (best guess from name/pronouns; N if unclear),
 "profile":{"facts":[{"category":str,"text":str,"source":"resume"|"ai_inference","verified":bool}],
  "gaps":[{"category":str,"question":str}],"insights":[str]},
 "slides":[{"t":"title","h":str,"s":str,"n":str},
  {"t":"timeline","h":str,"items":[[role,"company · dates — work performed"],...],"n":str},
  {"t":"skills","h":str,"items":[[skill,0-100],...],"n":str},
  {"t":"stats","h":str,"items":[[value,label],...],"n":str},   // only real figures from the resume; omit slide if none
  {"t":"list","h":str,"items":[[title,detail],...],"n":str},   // education/certs/projects
  {"t":"title","h":str,"s":str,"n":str}]}                       // closing
"n" is a friendly 20-40 second spoken script, third person, presented TO an audience.
Use tokens {name} {she} {her} {Her} instead of pronouns (e.g. "{Her} journey began...").
Keep the profile concise: up to 30 supported resume facts, clearly label AI interpretations as ai_inference and never mark them verified. Identify up to 4 useful story gaps as questions. Do not include raw resume text.
6-8 slides, max 4-5 items per slide, short on-slide text."""


INTERVIEW_SYSTEM = """You are Sassbyte's text-based career story editor, not a presenter or avatar.
Use only the candidate's uploaded-resume facts and what the candidate explicitly says in this conversation.
Do not turn guesses, questions, hypotheticals, or your own interpretation into candidate facts.
Ask focused questions about gaps such as ownership, challenge, outcome, metrics, leadership, and target role.
When the latest candidate message explicitly adds a factual detail or preference, propose it for approval.
Keep every proposed detail faithful to the candidate's words. Never invent metrics or rewrite an estimate as a fact.
Return only JSON matching:
{"reply":string,"proposals":[{"category":string,"text":string,"presentation_section":string}],"insights":[string]}
Use an empty proposals array unless the latest message clearly provides new candidate information.
The reply should acknowledge any new information, identify that it is a proposal (not yet saved), and ask at most one useful next question.
Keep the reply concise and friendly. Category should be one of: experience, achievement, leadership, project, skill, education, target_role, preference, other.
The presentation_section should be a short human-readable scene name."""

STORY_SYSTEM = """Create a career presentation from the candidate profile JSON provided by the user.
Return only JSON with {"name":string,"gender":"F"|"M"|"N","slides":[...]}.
Use only facts in the profile. Do not invent employers, dates, metrics, responsibilities, or skills.
Make a 5-8 scene narrative with concise visual text and a natural third-person narration script.
Each scene has: t ("title"|"timeline"|"skills"|"stats"|"list"), h, optional s, items (pairs for list/timeline/skills/stats), n, and fact_ids.
Include the exact text of every user_chat fact in at least one scene item and include its id in that scene's fact_ids.
Each user_chat fact is candidate-provided, not independently verified; present it as their account without adding claims.
Return concise scene content. Use tokens {name}, {she}, {her}, {Her} in narration when appropriate."""


def profile_from_deck(deck: dict) -> dict:
    facts = []
    present_categories = set()
    for slide in deck.get("slides", []):
        if not isinstance(slide, dict):
            continue
        slide_type = str(slide.get("t") or "list")
        heading = str(slide.get("h") or "Career detail")
        category = {
            "timeline": "experience",
            "skills": "skill",
            "stats": "achievement",
        }.get(slide_type, (
            "leadership" if "lead" in heading.lower()
            else "education" if any(word in heading.lower() for word in ("education", "certification"))
            else "achievement" if any(word in heading.lower() for word in ("impact", "achievement", "result"))
            else "project" if any(word in heading.lower() for word in ("project", "work"))
            else "other"
        ))
        present_categories.add(category)
        source = "ai_inference" if slide_type == "skills" else "resume"
        items = slide.get("items") if isinstance(slide.get("items"), list) else []
        if slide_type == "title":
            summary = str(slide.get("s") or "").strip()
            if summary:
                facts.append({
                    "id": uuid.uuid4().hex,
                    "category": "other",
                    "text": summary[:1000],
                    "source": "ai_inference",
                    "verified": False,
                })
        for item in items:
            if not isinstance(item, (list, tuple)) or not item:
                continue
            label = str(item[0]).strip()
            detail = str(item[1]).strip() if len(item) > 1 else ""
            fact_text = f"{label} — {detail}".strip(" —") if detail else label
            if not fact_text:
                continue
            facts.append({
                "id": uuid.uuid4().hex,
                "category": category,
                "text": fact_text[:1000],
                "source": source,
                "verified": source == "resume",
            })

    gaps = []
    if "achievement" not in present_categories:
        gaps.append({
            "category": "achievement",
            "question": "What changed because of your work? Any measurable result is helpful, but estimates are fine if you label them.",
        })
    if "leadership" not in present_categories:
        gaps.append({
            "category": "leadership",
            "question": "Did you guide people, influence a decision, or coordinate work across a team?",
        })
    if "experience" in present_categories:
        gaps.append({
            "category": "experience",
            "question": "Which role or project best represents the work you want people to remember?",
        })
    return {"facts": facts, "gaps": gaps, "insights": []}


def normalize_deck(deck: dict) -> dict:
    if not isinstance(deck, dict) or not isinstance(deck.get("slides"), list) or not deck["slides"]:
        raise HTTPException(502, "Storyboard response did not contain presentation scenes.")
    deck["name"] = str(deck.get("name") or "Your career story")[:160]
    deck["gender"] = deck.get("gender") if deck.get("gender") in {"F", "M", "N"} else "N"
    normalized_slides = []
    for slide in deck["slides"][:20]:
        if not isinstance(slide, dict):
            continue
        scene_type = slide.get("t")
        if scene_type not in {"title", "timeline", "skills", "stats", "list"}:
            scene_type = "list"
        items = slide.get("items")
        if not isinstance(items, list):
            items = []
        normalized_items = []
        for item in items[:12]:
            if not isinstance(item, (list, tuple)) or not item:
                continue
            label = str(item[0])[:300]
            detail = str(item[1])[:1000] if len(item) > 1 else ""
            if scene_type == "skills":
                try:
                    detail = str(max(0, min(100, float(detail))))
                except ValueError:
                    detail = "0"
            normalized_items.append([label, detail])
        fact_ids = slide.get("fact_ids")
        if not isinstance(fact_ids, list):
            fact_ids = []
        normalized_slides.append({
            "t": scene_type,
            "h": str(slide.get("h") or "Career story")[:240],
            "s": str(slide.get("s") or "")[:1000],
            "n": str(slide.get("n") or "")[:4000],
            "items": normalized_items,
            "fact_ids": [
                item for item in fact_ids[:50]
                if isinstance(item, str) and re.fullmatch(r"[a-f0-9]{32}", item)
            ],
        })
    deck["slides"] = normalized_slides
    if not deck["slides"]:
        raise HTTPException(502, "Storyboard response did not contain valid presentation scenes.")
    profile = deck.get("profile")
    if not isinstance(profile, dict) or not isinstance(profile.get("facts"), list):
        profile = profile_from_deck(deck)
    normalized_facts = []
    for fact in profile.get("facts", [])[:50]:
        if not isinstance(fact, dict):
            continue
        text = fact.get("text")
        if not isinstance(text, str) or not text.strip():
            continue
        source = fact.get("source")
        if source not in {"resume", "ai_inference"}:
            source = "ai_inference"
        normalized_facts.append({
            "id": uuid.uuid4().hex,
            "category": str(fact.get("category") or "other")[:60],
            "text": text.strip()[:1000],
            "source": source,
            "verified": source == "resume",
        })
    gaps = profile.get("gaps")
    if not isinstance(gaps, list):
        gaps = []
    normalized_gaps = [
        {
            "category": str(gap.get("category") or "other")[:60],
            "question": gap["question"].strip()[:500],
        }
        for gap in gaps[:6]
        if isinstance(gap, dict) and isinstance(gap.get("question"), str) and gap["question"].strip()
    ]
    insights = profile.get("insights")
    if not isinstance(insights, list):
        insights = []
    deck["profile"] = {
        "facts": normalized_facts,
        "gaps": normalized_gaps,
        "insights": [item.strip()[:500] for item in insights[:10] if isinstance(item, str) and item.strip()],
    }
    return deck


def llm_json(system: str, messages: list, max_tokens: int = 3000) -> dict:
    client_obj, client_error = get_client(use_env=True)
    if client_error:
        raise HTTPException(400, client_error)
    try:
        response = client_obj.messages.create(
            model=MODEL,
            max_tokens=max_tokens,
            system=system,
            messages=messages,
        )
    except Exception as exc:
        status_code, user_message = sanitize_error_response(str(exc))
        raise HTTPException(status_code, user_message) from exc
    raw = "".join(block.text for block in response.content if block.type == "text")
    match = re.search(r"\{.*\}", raw, re.S)
    if not match:
        raise HTTPException(502, "The story editor returned an invalid response. Please try again.")
    try:
        result = json.loads(match.group(0))
    except json.JSONDecodeError as exc:
        raise HTTPException(502, "The story editor returned an invalid response. Please try again.") from exc
    if not isinstance(result, dict):
        raise HTTPException(502, "The story editor returned an invalid response. Please try again.")
    return result


def parse(name: str, data: bytes) -> str:
    n = name.lower()
    if n.endswith(".pdf"):
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    if n.endswith(".docx"):
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text for c in r.cells))
        return "\n".join(parts)
    raise HTTPException(400, "Upload a PDF or DOCX file")

def storyboard(text: str, api_key: str = "", workspace_id: str = "", use_env: bool = False) -> dict:
    client_obj, client_error = get_client(api_key, workspace_id, use_env)
    if client_error:
        raise HTTPException(400, client_error)
    try:
        msg = client_obj.messages.create(model=MODEL, max_tokens=5000, system=SYSTEM,
                                        messages=[{"role": "user", "content": f"RESUME:\n{text[:MAX_CHARS]}"}])
    except Exception as exc:
        status_code, user_message = sanitize_error_response(str(exc))
        raise HTTPException(status_code, user_message) from exc
    raw = "".join(b.text for b in msg.content if b.type == "text")
    m = re.search(r"\{.*\}", raw, re.S)
    if not m:
        raise HTTPException(502, "Could not build storyboard")
    try:
        return json.loads(m.group(0))
    except json.JSONDecodeError as exc:
        raise HTTPException(502, "Storyboard JSON was invalid") from exc

@app.post("/api/upload")
async def upload(
    request: Request,
    file: UploadFile = File(...),
    api_key: str = Form(default=""),
    workspace_id: str = Form(default=""),
    use_env: bool = Form(default=False)
):
    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "File too large (max 5 MB)")
    text = parse(file.filename or "", data).strip()
    if len(text) < 100:
        raise HTTPException(422, "Could not read text from this file (scanned image PDFs are not supported yet)")
    try:
        deck = normalize_deck(storyboard(text, api_key=api_key, workspace_id=workspace_id, use_env=use_env))
    except HTTPException:
        raise
    except Exception as exc:
        status_code, user_message = sanitize_error_response(str(exc))
        raise HTTPException(status_code, user_message) from exc
    now = time.time()
    for k in [k for k, (e, _) in DECKS.items() if e < now]:
        DECKS.pop(k, None)
    did = uuid.uuid4().hex
    DECKS[did] = (now + TTL, deck)
    edit_token = secrets.token_urlsafe(32)
    share_url = f"{PUBLIC_BASE_URL}/share/{did}"
    deck["share_url"] = share_url
    persist_storyboard(did, deck)
    persist_edit_token(did, edit_token)
    return {"id": did, "deck": deck, "share_url": share_url, "edit_token": edit_token}


@app.post("/api/deck/{did}/interview")
async def interview(did: str, request: Request):
    token = request.headers.get("x-story-editor-token", "")
    if not can_edit_deck(did, token):
        raise HTTPException(403, "This story can no longer be edited from this browser.")
    deck = get_live_deck(did)
    if deck is None:
        raise HTTPException(404, "Storyboard expired or not found.")
    try:
        payload = await request.json()
    except Exception as exc:
        raise HTTPException(400, "A message is required.") from exc
    if not isinstance(payload, dict):
        raise HTTPException(400, "A message is required.")
    message = payload.get("message")
    if not isinstance(message, str) or not message.strip() or len(message) > 4000:
        raise HTTPException(422, "Enter a message of up to 4,000 characters.")

    history = payload.get("history", [])
    if not isinstance(history, list):
        raise HTTPException(422, "Conversation history is invalid.")
    safe_history = []
    for turn in history[-20:]:
        if not isinstance(turn, dict) or turn.get("role") not in {"user", "assistant"}:
            continue
        content = turn.get("content")
        if isinstance(content, str) and content.strip():
            safe_history.append({"role": turn["role"], "content": content[:4000]})

    profile = copy.deepcopy(deck.get("profile") or profile_from_deck(deck))
    system = INTERVIEW_SYSTEM + "\nCandidate facts (source labels are authoritative):\n" + json.dumps(
        profile.get("facts", [])[:80], ensure_ascii=False
    )
    result = llm_json(
        system,
        safe_history + [{"role": "user", "content": message.strip()}],
        max_tokens=1400,
    )
    reply = result.get("reply")
    if not isinstance(reply, str) or not reply.strip():
        raise HTTPException(502, "The story editor did not return a reply. Please try again.")
    allowed_categories = {
        "experience", "achievement", "leadership", "project", "skill",
        "education", "target_role", "preference", "other",
    }
    proposals = []
    for proposal in result.get("proposals", [])[:5] if isinstance(result.get("proposals"), list) else []:
        if not isinstance(proposal, dict):
            continue
        text = proposal.get("text")
        if not isinstance(text, str) or not text.strip():
            continue
        proposals.append({
            "category": proposal.get("category") if proposal.get("category") in allowed_categories else "other",
            "text": text.strip()[:1000],
            "presentation_section": str(proposal.get("presentation_section") or "Career story")[:120],
        })
    insights = [
        insight.strip()[:500]
        for insight in (result.get("insights") if isinstance(result.get("insights"), list) else [])
        if isinstance(insight, str) and insight.strip()
    ][:5]
    return {"reply": reply.strip()[:2000], "proposals": proposals, "insights": insights}


@app.post("/api/deck/{did}/story-facts")
async def add_story_fact(did: str, request: Request):
    token = request.headers.get("x-story-editor-token", "")
    if not can_edit_deck(did, token):
        raise HTTPException(403, "This story can no longer be edited from this browser.")
    deck = get_live_deck(did)
    if deck is None:
        raise HTTPException(404, "Storyboard expired or not found.")
    try:
        payload = await request.json()
    except Exception as exc:
        raise HTTPException(400, "A story detail is required.") from exc
    if not isinstance(payload, dict):
        raise HTTPException(400, "A story detail is required.")
    text = payload.get("text")
    if not isinstance(text, str) or not text.strip() or len(text) > 1000:
        raise HTTPException(422, "Story details must be between 1 and 1,000 characters.")
    categories = {
        "experience", "achievement", "leadership", "project", "skill",
        "education", "target_role", "preference", "other",
    }
    category = payload.get("category") if payload.get("category") in categories else "other"
    section = payload.get("presentation_section")
    if not isinstance(section, str):
        section = "Career story"
    section = section.strip()[:120] or "Career story"

    profile = copy.deepcopy(deck.get("profile") or profile_from_deck(deck))
    facts = profile.setdefault("facts", [])
    facts[:] = [fact for fact in facts if isinstance(fact, dict)]
    normalized_text = text.strip()
    duplicate = next(
        (fact for fact in facts if fact.get("source") == "user_chat" and fact.get("text", "").casefold() == normalized_text.casefold()),
        None,
    )
    if duplicate:
        return {"deck": deck, "fact": duplicate, "message": "That detail is already in your story."}
    if sum(1 for item in facts if item.get("source") == "user_chat") >= 30:
        raise HTTPException(409, "This story already has 30 candidate-added details.")

    fact = {
        "id": uuid.uuid4().hex,
        "category": category,
        "text": normalized_text,
        "source": "user_chat",
        "verified": True,
        "presentation_section": section,
    }
    facts.append(fact)
    profile["gaps"] = [
        gap for gap in profile.get("gaps", [])
        if not isinstance(gap, dict) or gap.get("category") != category
    ]
    updated = llm_json(
        STORY_SYSTEM,
        [{"role": "user", "content": json.dumps(profile, ensure_ascii=False)}],
        max_tokens=4000,
    )
    updated["name"] = deck.get("name", "Your career story")
    updated["gender"] = deck.get("gender", "N")
    updated = normalize_deck(updated)
    updated["profile"] = profile

    for user_fact in (item for item in facts if item.get("source") == "user_chat"):
        exact_fact_present = False
        for slide in updated["slides"]:
            visible_content = [
                str(slide.get("h") or ""),
                str(slide.get("s") or ""),
                *(str(value) for item in slide.get("items", []) for value in item),
            ]
            slide_text = json.dumps(visible_content, ensure_ascii=False)
            if user_fact["text"] in slide_text:
                fact_ids = slide.get("fact_ids")
                if not isinstance(fact_ids, list):
                    fact_ids = []
                if user_fact["id"] not in fact_ids:
                    fact_ids.append(user_fact["id"])
                slide["fact_ids"] = fact_ids
                exact_fact_present = True
                break
        if not exact_fact_present:
            updated["slides"].append({
                "t": "list",
                "h": user_fact.get("presentation_section") or "Added by you",
                "items": [[user_fact["category"].replace("_", " ").title(), user_fact["text"]]],
                "n": f"This detail was added by {updated['name']} during the story interview.",
                "fact_ids": [user_fact["id"]],
            })

    updated["share_url"] = f"{PUBLIC_BASE_URL}/share/{did}"
    DECKS[did] = (time.time() + TTL, updated)
    persist_storyboard(did, updated)
    return {"deck": updated, "fact": fact, "message": "Added to your story and presentation."}

@app.get("/api/deck/{did}")
def get_deck(did: str):
    deck = get_live_deck(did)
    if deck is None:
        raise HTTPException(404, "Deck expired or not found")
    return deck


@app.get("/api/decks")
def list_decks():
    decks = []
    for entry in sorted(os.listdir(PERSIST_BASE)):
        deck_dir = os.path.join(PERSIST_BASE, entry)
        if not os.path.isdir(deck_dir):
            continue
        deck_path = os.path.join(deck_dir, "deck.json")
        if not os.path.exists(deck_path):
            continue
        try:
            with open(deck_path, "r", encoding="utf-8") as f:
                deck = json.load(f)
            decks.append({"id": entry, "name": (deck.get("name") or "Resume Storyboard").strip() or "Resume Storyboard"})
        except Exception:
            continue
    return {"decks": decks}

PRON = {"F": ("she", "her", "She"), "M": ("he", "his", "He"), "N": ("they", "their", "They")}
def fill(s, deck, g):
    a, b, c = PRON.get(g, PRON["N"])
    return s.replace("{name}", deck.get("name", "")).replace("{she}", a).replace("{her}", b).replace("{Her}", c)

@app.get("/api/deck/{did}/pptx")
def export(did: str, g: str = "N"):
    deck = get_live_deck(did)
    if deck is None:
        raise HTTPException(404, "Deck expired or not found")
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    for s in deck["slides"]:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        tb = sl.shapes.add_textbox(Inches(.8), Inches(.6), Inches(11.7), Inches(1.2)).text_frame
        tb.text = fill(s["h"], deck, g); tb.paragraphs[0].font.size = Pt(40); tb.paragraphs[0].font.bold = True
        body = sl.shapes.add_textbox(Inches(.8), Inches(2), Inches(11.7), Inches(4.8)).text_frame
        body.word_wrap = True
        lines = [fill(s.get("s", ""), deck, g)] if s["t"] == "title" else [
            f"{a}  —  {b}{'%' if s['t']=='skills' else ''}" for a, b in s.get("items", [])]
        for k, ln in enumerate(lines):
            p = body.paragraphs[0] if k == 0 else body.add_paragraph()
            p.text = ln; p.font.size = Pt(24); p.space_after = Pt(12)
        sl.notes_slide.notes_text_frame.text = fill(s.get("n", ""), deck, g)
    buf = io.BytesIO(); prs.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                             headers={"Content-Disposition": 'attachment; filename="resume-presentation.pptx"'})

@app.get("/api/share/{did}")
def get_shared_deck(did: str):
    deck = load_persisted_storyboard(did)
    if not deck:
        raise HTTPException(404, "Shared storyboard not found")
    return deck

@app.get("/share/{did}")
def share_storyboard_page(did: str):
    deck = load_persisted_storyboard(did)
    if not deck:
        raise HTTPException(404, "Shared storyboard not found")
    return HTMLResponse(render_share_page(did, deck))

@app.get("/rp")
@app.get("/rp/")
async def rp_index():
    return FileResponse("index.html")

@app.get("/resume-presenter")
@app.get("/resume-presenter/")
async def resume_presenter_index():
    return FileResponse("index.html")

@app.get("/rp/static")
@app.get("/rp/static/")
async def rp_static_index():
    return FileResponse("static/index.html")

@app.get("/rp/static/{path:path}")
async def rp_static_file(path: str):
    safe_path = os.path.normpath(os.path.join("static", path))
    if not safe_path.startswith("static"):
        raise HTTPException(403, "Forbidden")
    if os.path.isdir(safe_path):
        safe_path = os.path.join(safe_path, "index.html")
    if not os.path.exists(safe_path):
        raise HTTPException(404, "Not found")
    return FileResponse(safe_path)

@app.get("/resume-presenter/static")
@app.get("/resume-presenter/static/")
async def resume_presenter_static_index():
    return FileResponse("static/index.html")

@app.get("/resume-presenter/static/{path:path}")
async def resume_presenter_static_file(path: str):
    safe_path = os.path.normpath(os.path.join("static", path))
    if not safe_path.startswith("static"):
        raise HTTPException(403, "Forbidden")
    if os.path.isdir(safe_path):
        safe_path = os.path.join(safe_path, "index.html")
    if not os.path.exists(safe_path):
        raise HTTPException(404, "Not found")
    return FileResponse(safe_path)

app.mount("/", StaticFiles(directory="static", html=True), name="static")
